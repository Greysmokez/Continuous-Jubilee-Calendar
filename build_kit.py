#!/usr/bin/env python3
"""
Build every public copy of the CJC Research Kit from ONE master file.

  kit-source/kit-master.txt   <- the only kit file you ever edit
  kit-source/si-header.txt    <- SI-edition front matter (edit rarely)
  kit-source/kit-rules.json   <- consistency rules checked before anything is written

Outputs (same filenames and paths as before, so no links change):
  CJCKit.txt                  SI header + master text     (this repo)
  CJC_AI_Research_Kit.pdf     PDF built from the master   (this repo)
  <website-out>/kit.txt       exact copy of the master    (pushed to CJC-Website by the workflow)

Run locally:
  python scripts/build_kit.py --website-out build-website
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "kit-source"
MASTER = SRC / "kit-master.txt"
HEADER = SRC / "si-header.txt"
RULES = SRC / "kit-rules.json"

VERSION_RE = re.compile(r"^Version (\d+(?:\.\d+)*) — (.+?\d{4})\.")
BLUE = RGBColor(0x4F, 0x81, 0xBD)

# Paragraph openings shown in bold in the PDF (the .txt files are never changed).
LEAD_PATTERNS = [
    re.compile(r"^(Version [^.]*\.[^.]*\.)"),
    re.compile(r"^(Anchor \d+ — [^.]*\.)"),
    re.compile(r"^(Mid-point, defined once, up front:.*?\.)(?= |$)"),
    re.compile(r"^(Important structural finding:)"),
    re.compile(r"^(The law\.|The imperial coins\.|The provincial coins\.|How to present this\.|"
               r"How to apply this when answering\.|Three cautions\.)"),
    re.compile(r"^(First use — [^.]*\.[^.]*\.)"),
    re.compile(r"^(Second use — [^.]*\.)"),
    re.compile(r"^(\d+\. [^?]{5,200}\?)"),
]


def fail(msg: str) -> None:
    print(f"::error::{msg}")
    sys.exit(1)


# ---------------------------------------------------------------- checks
def check_master(text: str) -> tuple[str, str]:
    lines = text.split("\n")
    if len(lines) < 5 or not lines[0].strip():
        fail("kit-master.txt must start with the kit title on line 1.")
    m = VERSION_RE.match(lines[2].strip()) if len(lines) > 2 else None
    if not m:
        fail('Line 3 of kit-master.txt must be the version line, e.g. '
             '"Version 12.16 — October 1, 2026. Full revision history ..."')
    if "How to Use This Document" not in lines:
        fail('kit-master.txt must contain the heading "How to Use This Document".')

    if RULES.exists():
        rules = json.loads(RULES.read_text(encoding="utf-8"))
        problems = []
        for r in rules.get("must_include", []):
            if r["text"] not in text:
                problems.append(f'MISSING "{r["text"]}" ({r.get("why", "")})')
        for r in rules.get("must_not_include", []):
            if r["text"] in text:
                problems.append(f'RETIRED PHRASE PRESENT "{r["text"]}" ({r.get("why", "")})')
        if problems:
            for p in problems:
                print(f"::error::{p}")
            fail(f"{len(problems)} consistency rule(s) failed; nothing was published.")
    return m.group(1), m.group(2)


# ---------------------------------------------------------------- text outputs
def build_si_edition(master: str) -> str:
    header = HEADER.read_text(encoding="utf-8").rstrip("\n")
    lines = master.split("\n")
    version_line = lines[2].strip()
    underlying = "Underlying kit text v" + version_line[1:]
    body = "\n".join(lines[lines.index("How to Use This Document"):])
    return f"{header}\n\n{underlying}\n\n{body}"


# ---------------------------------------------------------------- PDF output
def is_heading(line: str) -> bool:
    s = line.strip()
    return (
        0 < len(s) <= 110
        and not s.endswith((".", ":", "?", ";", ","))
        and not s.startswith(("• ", "Version "))
        and not re.search(r"[=≤×]", s)
        and not re.match(r"^\d+\.", s)
    )


def is_formula(line: str) -> bool:
    s = line.strip()
    return len(s) <= 110 and (bool(re.search(r"[=≤]", s)) or s in {"From B:"}
                              or s.startswith("Apply the main correction"))


def add_runs(par, text: str, bold: bool = False) -> None:
    for part in re.split(r"(\*[^*\n]+\*)", text):
        if not part:
            continue
        italic = len(part) > 2 and part.startswith("*") and part.endswith("*")
        run = par.add_run(part[1:-1] if italic else part)
        run.italic = italic
        run.bold = bold


def build_docx(master: str, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Cambria", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name, size in (("Title", 18), ("Heading 1", 13)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.font.color.rgb = BLUE
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        # remove theme border/underline Word adds to Title
        pPr = st.element.get_or_add_pPr()
        for child in list(pPr):
            if child.tag.endswith("pBdr"):
                pPr.remove(child)

    first = True
    for raw in master.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if first:
            doc.add_paragraph(line, style="Title")
            first = False
            continue
        if line.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
            continue
        if is_formula(line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, line)
            continue
        if is_heading(line):
            doc.add_paragraph(line, style="Heading 1")
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lead = next((m.group(1) for pat in LEAD_PATTERNS if (m := pat.match(line))), None)
        if lead:
            add_runs(p, lead, bold=True)
            add_runs(p, line[len(lead):])
        else:
            add_runs(p, line)
    doc.save(str(out))


def docx_to_pdf(docx_path: Path, out_dir: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        fail("LibreOffice (soffice) is required to build the PDF.")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = out_dir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        fail("PDF conversion produced no file.")
    return pdf


# ---------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--website-out", type=Path, required=True,
                    help="Folder to write kit.txt for the CJC-Website repo (kept outside this repo in CI).")
    ap.add_argument("--skip-pdf", action="store_true", help="Build the .txt files only.")
    args = ap.parse_args()

    master = MASTER.read_text(encoding="utf-8").replace("\r\n", "\n")
    version, date = check_master(master)
    print(f"Master kit: version {version} — {date}")

    (ROOT / "CJCKit.txt").write_text(build_si_edition(master), encoding="utf-8")
    print("Wrote CJCKit.txt")

    args.website_out.mkdir(parents=True, exist_ok=True)
    (args.website_out / "kit.txt").write_text(master, encoding="utf-8")
    print(f"Wrote {args.website_out / 'kit.txt'}")

    if not args.skip_pdf:
        with tempfile.TemporaryDirectory() as tmp:  # temp files never land in the repo
            tmp = Path(tmp)
            docx_path = tmp / "CJC_AI_Research_Kit.docx"
            build_docx(master, docx_path)
            pdf = docx_to_pdf(docx_path, tmp)
            shutil.copyfile(pdf, ROOT / "CJC_AI_Research_Kit.pdf")
        print("Wrote CJC_AI_Research_Kit.pdf")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
