#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.workbook.protection import WorkbookProtection
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import Color
from reportlab.pdfgen import canvas


NOTICE_FOOTER = "© 2026 Chip Welsh. All Rights Reserved. Continuous Jubilee Calendar™ and CJC™ are trademarks of Chip Welsh."
NOTICE_SHORT = "Continuous Jubilee Calendar™ / CJC™ — © 2026 Chip Welsh"
ABOUT_LINES = [
    "© 2026 Chip Welsh. All Rights Reserved.",
    "Continuous Jubilee Calendar™ and CJC™ are trademarks of Chip Welsh.",
    "Terms: see legal/legal-notice.txt (Effective Date: September 22, 2026).",
    "Notice scope: these labels communicate ownership claims and handling expectations; they do not create legal rights beyond applicable law.",
]
WATERMARK_COLOR = (140, 140, 140)
WATERMARK_OPACITY = 0.08


def ensure_dir(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def relative_target(source_root: Path, output_root: Path, file_path: Path) -> Path:
    rel = file_path.relative_to(source_root)
    return output_root / rel


def add_docx_notice(source: Path, target: Path) -> None:
    doc = Document(str(source))
    for section in doc.sections:
        footer = section.footer
        if any(NOTICE_FOOTER in p.text for p in footer.paragraphs):
            continue
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        if paragraph.text.strip():
            paragraph.add_run("  ")
        run = paragraph.add_run(NOTICE_FOOTER)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(130, 130, 130)
    ensure_dir(target)
    doc.save(str(target))


def _overlay_stream(page_width: float, page_height: float, text: str, top_text: str) -> bytes:
    from io import BytesIO

    stream = BytesIO()
    c = canvas.Canvas(stream, pagesize=(page_width, page_height))
    c.setFillColor(Color(0.55, 0.55, 0.55, alpha=WATERMARK_OPACITY))
    c.setFont("Helvetica", 8)
    c.drawCentredString(page_width / 2, 14, text)
    c.drawString(12, page_height - 14, top_text)
    c.save()
    stream.seek(0)
    return stream.read()


def add_pdf_notice(source: Path, target: Path) -> None:
    from io import BytesIO

    reader = PdfReader(str(source))
    writer = PdfWriter()
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_bytes = _overlay_stream(width, height, NOTICE_FOOTER, NOTICE_SHORT)
        overlay_page = PdfReader(BytesIO(overlay_bytes)).pages[0]
        page.merge_page(overlay_page, over=True)
        writer.add_page(page)
    ensure_dir(target)
    with target.open("wb") as fh:
        writer.write(fh)


def add_xlsx_about_sheet(source: Path, target: Path, workbook_structure_password: str | None = None) -> None:
    keep_vba = source.suffix.lower() == ".xlsm"
    wb = load_workbook(filename=str(source), keep_vba=keep_vba)
    if "About" in wb.sheetnames:
        del wb["About"]
    about = wb.create_sheet("About", 0)
    for idx, line in enumerate(ABOUT_LINES, start=1):
        cell = about.cell(row=idx, column=1, value=line)
        cell.font = Font(color="808080", size=10, bold=(idx == 1))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    about.column_dimensions["A"].width = 120
    about.protection.sheet = True
    if workbook_structure_password:
        wb.security = WorkbookProtection(lockStructure=True, workbookPassword=workbook_structure_password)
    ensure_dir(target)
    wb.save(str(target))


def load_watermark_font(watermark_font_path: str | None) -> ImageFont.ImageFont:
    if watermark_font_path:
        return ImageFont.truetype(watermark_font_path, 12)
    try:
        return ImageFont.truetype("DejaVuSans.ttf", 12)
    except OSError:
        return ImageFont.load_default()


def add_image_margin_watermark(source: Path, target: Path, watermark_font_path: str | None = None) -> None:
    with Image.open(source) as original:
        image = original.convert("RGBA")
        margin = max(16, int(min(image.width, image.height) * 0.04))
        supports_alpha = target.suffix.lower() not in {".jpg", ".jpeg"}
        background = (255, 255, 255, 0) if supports_alpha else (255, 255, 255, 255)
        canvas_img = Image.new(
            "RGBA",
            (image.width + margin * 2, image.height + margin * 2),
            background,
        )
        canvas_img.paste(image, (margin, margin))

        layer = Image.new("RGBA", canvas_img.size, (255, 255, 255, 0))
        draw = ImageDraw.Draw(layer)
        font = load_watermark_font(watermark_font_path)
        alpha = int(255 * WATERMARK_OPACITY)
        fill = (*WATERMARK_COLOR, alpha)
        text = NOTICE_SHORT
        bbox = draw.textbbox((0, 0), text, font=font)
        text_w = bbox[2] - bbox[0]
        text_h = bbox[3] - bbox[1]
        draw.text(((canvas_img.width - text_w) / 2, 4), text, fill=fill, font=font)
        draw.text(
            ((canvas_img.width - text_w) / 2, canvas_img.height - text_h - 4),
            text,
            fill=fill,
            font=font,
        )

        merged = Image.alpha_composite(canvas_img, layer)
        target_suffix = target.suffix.lower()
        if target_suffix in {".jpg", ".jpeg"}:
            merged = merged.convert("RGB")
        ensure_dir(target)
        merged.save(target)


def copy_textual_notice(source: Path, target: Path) -> None:
    ensure_dir(target)
    shutil.copy2(source, target)


def process(
    source_root: Path,
    output_root: Path,
    workbook_structure_password: str | None = None,
    watermark_font_path: str | None = None,
) -> dict:
    supported = {".pdf", ".docx", ".xlsx", ".xlsm", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}
    skipped = []
    processed = []
    source_root = source_root.resolve()
    output_root = output_root.resolve()
    output_rel_parts = None
    try:
        output_rel_parts = output_root.relative_to(source_root).parts
    except ValueError:
        output_rel_parts = None

    for file_path in source_root.rglob("*"):
        if not file_path.is_file():
            continue
        rel_file = file_path.relative_to(source_root)
        if output_rel_parts and rel_file.parts[: len(output_rel_parts)] == output_rel_parts:
            continue
        resolved_file = file_path.resolve()
        if output_root in resolved_file.parents or ".git" in file_path.parts:
            continue
        target = relative_target(source_root, output_root, file_path)
        suffix = file_path.suffix.lower()
        try:
            if suffix == ".pdf":
                add_pdf_notice(file_path, target)
                processed.append(str(file_path.relative_to(source_root)))
            elif suffix == ".docx":
                add_docx_notice(file_path, target)
                processed.append(str(file_path.relative_to(source_root)))
            elif suffix in {".xlsx", ".xlsm"}:
                add_xlsx_about_sheet(
                    file_path,
                    target,
                    workbook_structure_password=workbook_structure_password,
                )
                processed.append(str(file_path.relative_to(source_root)))
            elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
                add_image_margin_watermark(file_path, target, watermark_font_path=watermark_font_path)
                processed.append(str(file_path.relative_to(source_root)))
            elif suffix == ".txt" and file_path.name == "legal-notice.txt":
                copy_textual_notice(file_path, target)
                processed.append(str(file_path.relative_to(source_root)))
            elif suffix not in supported:
                skipped.append(str(file_path.relative_to(source_root)))
        except Exception as exc:  # pragma: no cover - operational reporting
            skipped.append(f"{file_path.relative_to(source_root)} ({exc})")

    output_root.mkdir(parents=True, exist_ok=True)
    report = {
        "effective_date": "September 22, 2026",
        "watermark_opacity": WATERMARK_OPACITY,
        "watermark_color": WATERMARK_COLOR,
        "margin_only": True,
        "workbook_structure_locked": bool(workbook_structure_password),
        "processed_files": sorted(processed),
        "skipped_files": sorted(skipped),
    }
    (output_root / "processing-report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply legal labels/watermarks to source assets without editing originals.")
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("."),
        help="Source directory containing original documents/assets.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("protected-assets"),
        help="Output directory where protected copies are written.",
    )
    parser.add_argument(
        "--workbook-structure-password",
        default=None,
        help="Optional workbook structure password. If omitted, only the About worksheet itself is locked.",
    )
    parser.add_argument(
        "--watermark-font-path",
        default=None,
        help="Optional TTF font path for deterministic image watermark rendering.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_root = args.output.resolve()
    report = process(
        args.source.resolve(),
        output_root,
        workbook_structure_password=args.workbook_structure_password,
        watermark_font_path=args.watermark_font_path,
    )
    print(f"Processed files: {len(report['processed_files'])}")
    print(f"Skipped files: {len(report['skipped_files'])}")
    print(f"Report: {output_root / 'processing-report.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
